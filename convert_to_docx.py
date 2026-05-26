import re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    """Set cell background color."""
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color)
    shading_elm.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_border(cell, **kwargs):
    """Set cell borders."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge, val in kwargs.items():
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), val.get('val', 'single'))
        element.set(qn('w:sz'), val.get('sz', '4'))
        element.set(qn('w:color'), val.get('color', '000000'))
        element.set(qn('w:space'), '0')
        tcBorders.append(element)
    tcPr.append(tcBorders)

def add_styled_table(doc, headers, rows, col_widths=None):
    """Add a professionally styled table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(header)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(cell, '4F46E5')

    # Data rows
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            if r % 2 == 1:
                set_cell_shading(cell, 'F3F4F6')

    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(width)

    doc.add_paragraph('')  # spacing
    return table

def add_heading_with_style(doc, text, level=1):
    """Add a heading with custom formatting."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        if level == 1:
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
        elif level == 2:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x37, 0x3E, 0x99)
        elif level == 3:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x43, 0x3E, 0x9E)
    return heading

def add_body_text(doc, text):
    """Add body text with consistent formatting."""
    if not text.strip():
        return
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_bullet(doc, text, level=0):
    """Add a bullet point."""
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    if level > 0:
        p.paragraph_format.left_indent = Cm(1.5 * (level + 1))
    return p

def add_code_block(doc, text):
    """Add a code-styled paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
    set_cell_shading_paragraph(p, 'F0F4FF')
    return p

def set_cell_shading_paragraph(paragraph, color):
    """Set paragraph background shading."""
    pPr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    pPr.append(shading)

def parse_markdown_to_docx(md_path, docx_path, title=None):
    """Convert a markdown file to a styled .docx file."""
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    lines = content.split('\n')
    i = 0
    in_table = False
    table_headers = []
    table_rows = []
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # Skip empty lines
        if not line.strip():
            if in_code_block:
                code_lines.append('')
            elif in_table and table_rows:
                # End of table
                add_styled_table(doc, table_headers, table_rows)
                table_headers = []
                table_rows = []
                in_table = False
            i += 1
            continue

        # Code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                for cl in code_lines:
                    add_code_block(doc, cl)
                doc.add_paragraph('')
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            add_heading_with_style(doc, line[2:], level=1)
        elif line.startswith('## '):
            add_heading_with_style(doc, line[2:], level=2)
        elif line.startswith('### '):
            add_heading_with_style(doc, line[2:], level=3)
        elif line.startswith('#### '):
            add_heading_with_style(doc, line[2:], level=3)

        # Horizontal rule
        elif line.strip() == '---':
            doc.add_paragraph('─' * 60)
            doc.add_paragraph('')

        # Bullet points
        elif line.strip().startswith('- '):
            text = line.strip()[2:]
            # Handle bold within bullets
            text = re.sub(r'\*\*(.+?)\*\*', lambda m: m.group(1), text)
            text = re.sub(r'`(.+?)`', lambda m: m.group(1), text)
            add_bullet(doc, text)

        # Blockquotes (used for time/speaker indicators)
        elif line.strip().startswith('> '):
            text = line.strip()[2:]
            text = re.sub(r'\*\*(.+?)\*\*', lambda m: m.group(1), text)
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.size = Pt(10)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        # Table rows
        elif line.strip().startswith('|'):
            # Skip separator rows
            if re.match(r'^\|[\s\-:]+\|', line):
                i += 1
                continue

            cells = [c.strip() for c in line.strip('|').split('|')]

            # Detect if this is a header or data row
            if not in_table:
                # Check if next line is a separator
                if i + 1 < len(lines) and re.match(r'^\|[\s\-:]+\|', lines[i + 1]):
                    table_headers = [re.sub(r'\*\*(.+?)\*\*', r'\1', c) for c in cells]
                    in_table = True
                    i += 1
                else:
                    in_table = True
                    table_headers = [re.sub(r'\*\*(.+?)\*\*', r'\1', c) for c in cells]
            else:
                cleaned = [re.sub(r'\*\*(.+?)\*\*', r'\1', c) for c in cells]
                cleaned = [re.sub(r'`(.+?)`', r'\1', c) for c in cleaned]
                table_rows.append(cleaned)
        else:
            # Regular text
            text = line.strip()
            # Handle bold
            text = re.sub(r'\*\*(.+?)\*\*', lambda m: m.group(1), text)
            # Handle inline code
            text = re.sub(r'`(.+?)`', lambda m: m.group(1), text)
            add_body_text(doc, text)

        i += 1

    # Flush any remaining table
    if in_table and table_headers:
        add_styled_table(doc, table_headers, table_rows)

    doc.save(docx_path)
    print(f'Created: {docx_path}')

if __name__ == '__main__':
    base = r'C:\Users\HP\Documents\IP project\expense-tracker'

    parse_markdown_to_docx(
        f'{base}\\WORKLOAD_DIVISION.md',
        f'{base}\\Workload_Division.docx',
        'Workload Division — Spendly Expense Tracker'
    )

    parse_markdown_to_docx(
        f'{base}\\PRESENTATION_SCRIPT_MEMBER1.md',
        f'{base}\\Presentation_Script_Member1.docx',
        'Presentation Script — Member 1'
    )

    parse_markdown_to_docx(
        f'{base}\\PRESENTATION_SCRIPT_MEMBER2.md',
        f'{base}\\Presentation_Script_Member2.docx',
        'Presentation Script — Member 2'
    )

    print('\nAll .docx files created successfully.')
